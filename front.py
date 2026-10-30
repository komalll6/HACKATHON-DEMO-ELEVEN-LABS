import streamlit as st
from docx import Document
import io

def generate_resume_docx(name, email, role, skills, experience):
    doc = Document()
    
    # Header
    doc.add_heading(name, level=0)
    doc.add_paragraph(f"Email: {email} | Role: {role}")
    
    # Skills Section
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(skills)
    
    # Experience Section
    doc.add_heading("Work Experience / Projects", level=1)
    doc.add_paragraph(experience)
    
    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Streamlit UI Component
st.subheader("📄 CareerShield AI Resume Generator")

col1, col2 = st.columns(2)
with col1:
    name = st.text_input("Full Name")
    email = st.text_input("Email Address")
    role = st.text_input("Target Job Role", value="MERN Stack Developer")

with col2:
    skills = st.text_area("Skills (Comma Separated)", "React.js, Node.js, Express, MongoDB, Tailwind CSS")
    experience = st.text_area("Experience / Key Projects", "Developed CareerShield AI voice interview agent...")

if st.button("Generate & Download Resume"):
    if name and email:
        resume_file = generate_resume_docx(name, email, role, skills, experience)
        st.download_button(
            label="📥 Download Resume (.docx)",
            data=resume_file,
            file_name=f"{name.replace(' ', '_')}_Resume.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    else:
        st.error("Please fill in Name and Email!")